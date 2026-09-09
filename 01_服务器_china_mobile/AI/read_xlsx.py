"""
中间代码：读取 任务.xlsx 并输出全部内容
用于解析课题任务 Excel 文件
"""
import openpyxl
import sys
import io

# 确保 UTF-8 输出，避免 GBK 编码报错
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX_PATH = r"e:\china_mobile\任务.xlsx"

def read_xlsx(path: str):
    wb = openpyxl.load_workbook(path)
    print(f"文件: {path}")
    print(f"Sheet 数量: {len(wb.sheetnames)}, 名称: {wb.sheetnames}")
    print("=" * 100)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n>>> Sheet: {sheet_name}")
        print(f"    维度: {ws.dimensions}, 行数: {ws.max_row}, 列数: {ws.max_column}\n")

        for i, row in enumerate(ws.iter_rows(
            min_row=1, max_row=ws.max_row,
            max_col=ws.max_column, values_only=False
        ), start=1):
            print(f"--- 第 {i} 行 ---")
            for j, cell in enumerate(row):
                col_letter = chr(ord('A') + j)
                val = cell.value if cell.value is not None else "(空)"
                # 截断过长内容
                val_str = str(val)
                if len(val_str) > 300:
                    val_str = val_str[:300] + f"\n    ... [共 {len(val_str)} 字符，已截断]"
                print(f"  [{col_letter}] {val_str}")
            print()

def extract_task_detail(path: str):
    """提取子课题一和子课题二的详细要求（第4行 B/D 列）"""
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]

    print("=" * 100)
    print("子课题一 详细要求 (Row 4, Col B)")
    print("=" * 100)
    b4 = ws['B4'].value
    if b4:
        print(b4)

    print("\n" + "=" * 100)
    print("子课题二 详细要求 (Row 4, Col D)")
    print("=" * 100)
    d4 = ws['D4'].value
    if d4:
        print(d4)

    print("\n" + "=" * 100)
    print("周报格式参考 (Row 4, Col C)")
    print("=" * 100)
    c4 = ws['C4'].value
    if c4:
        print(c4)

if __name__ == "__main__":
    print("【完整读取】")
    read_xlsx(XLSX_PATH)
    print("\n\n")
    extract_task_detail(XLSX_PATH)
